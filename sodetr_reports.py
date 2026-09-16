"""CPU-only experiment collection and JSON/CSV/Markdown/Excel reports.

Ratios are stored in JSON/CSV; human-facing tables use percent and deltas use pp.
No checkpoint is unpickled while scanning experiments.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import os
import tempfile
from datetime import datetime, timezone
from pathlib import Path

import yaml

from sodetr_formal_coco import COCO_STAT_NAMES, enrich_saved_evaluation

ROOT = Path(__file__).resolve().parent
MODES = {
    "fixed": ("Q0", "基线：固定 Expanded-IoU"),
    "adaptive-quality": ("Q1", "仅查询质量监督使用尺度自适应 Expanded-IoU"),
    "adaptive-regression": ("Q2", "仅回归损失使用尺度自适应 Expanded-IoU"),
    "adaptive-both": ("Q3", "查询质量监督与回归损失均使用尺度自适应 Expanded-IoU"),
}
COMPARE_KEYS = ("seed", "imgsz", "batch", "epochs", "patience", "optimizer", "lr0", "lrf",
                "weight_decay", "amp", "deterministic", "pretrained", "mixup", "close_mosaic",
                "warmup_epochs", "mosaic", "fraction", "expanded_iou_fixed_ratio",
                "expanded_iou_alpha", "expanded_iou_tau", "expanded_iou_min_ratio", "expanded_iou_max_ratio")


def write_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(dir=path.parent, suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            json.dump(data, stream, ensure_ascii=False, indent=2, allow_nan=False)
            stream.write("\n")
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)


def read_json(path: Path) -> dict:
    # A run may have created the metadata file before being interrupted.
    # Treat an empty file like a missing optional metadata file so report
    # finalization can still proceed and rebuild the metadata from state.
    if not path.is_file() or path.stat().st_size == 0:
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def number(value):
    try:
        value = float(value)
        return value if math.isfinite(value) else None
    except (TypeError, ValueError):
        return None


def read_curves(path: Path) -> list[dict]:
    if not path.is_file():
        return []
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return _curves(stream)


def _curves(stream):
    rows = []
    for raw in csv.DictReader(stream):
        row = {k.strip(): number(v) for k, v in raw.items() if k is not None}
        if row.get("epoch") is not None:
            rows.append(row)
    return rows


def collect_run(run: Path) -> dict:
    config = yaml.safe_load((run / "args.yaml").read_text(encoding="utf-8")) if (run / "args.yaml").exists() else {}
    config = config or {}
    meta = read_json(run / "experiment.json")
    state = read_json(run / "training_state.json")
    evaluation = read_json(run / "evaluation_state.json")
    coco_path = run / "formal_coco/best/coco_metrics.json"
    coco = read_json(coco_path)
    curves = read_curves(run / "results.csv")
    mode = config.get("expanded_iou_mode", "unknown")
    experiment, description = MODES.get(mode, (run.name, "未提供实验说明"))
    warnings = []
    # A new training session or failed re-evaluation makes older metrics stale.
    stale = evaluation.get("status") in ("pending", "running", "failed", "skipped") or state.get("status") in ("running", "failed", "interrupted")
    weight_path = run / "weights/best.pt"
    if coco.get("weights_mtime_ns") is not None and weight_path.exists():
        stale = stale or weight_path.stat().st_mtime_ns != coco["weights_mtime_ns"]
    if coco and stale:
        warnings.append("已有 COCO 结果已过期，本次不纳入比较")
    metrics = {k: number(v) if number(v) is not None and number(v) >= 0 else None
               for k, v in coco.get("metrics", {}).items()} if not stale else {}
    best_epoch = state.get("best_epoch")
    best_source = state.get("best_epoch_source", "trainer") if best_epoch else "unknown"
    valid = [r for r in curves if r.get("metrics/mAP50(B)") is not None and r.get("metrics/mAP50-95(B)") is not None]
    if best_epoch is None and valid:
        best = max(valid, key=lambda r: (.1 * r["metrics/mAP50(B)"] + .9 * r["metrics/mAP50-95(B)"], r["epoch"]))
        best_epoch = int(best["epoch"])
        best_source = "estimated_from_rounded_csv_fitness"
        warnings.append("最佳轮次及 P/R 按 CSV 中 0.1×mAP50+0.9×mAP50-95 估算")
    best_row = next((r for r in reversed(curves) if r["epoch"] == best_epoch), {})
    if not coco:
        warnings.append("尚无正式 COCOeval 结果")
    if state.get("status", "unknown") == "unknown":
        warnings.append("历史实验未记录训练状态、停止原因和耗时")
    if coco and not coco.get("per_class"):
        warnings.append("缺少每类别 AP，可使用 --backfill-per-class 补算")
    file_paths = {"run": str(run.resolve()), "best_weights": str(run / "weights/best.pt"),
                  "results_csv": str(run / "results.csv"), "coco_metrics": str(coco_path),
                  "args": str(run / "args.yaml")}
    model = meta.get("model") or state.get("model") or config.get("model")
    # A resume checkpoint path does not reveal the model architecture.
    if model and str(model).endswith(".pt") and not state.get("model") and not meta.get("model"):
        warnings.append("模型配置仅记录为 checkpoint 路径，无法确认与其他实验架构一致")
    return {
        "run": run.name, "experiment": meta.get("experiment", experiment),
        "description": meta.get("description") or description, "mode": mode,
        "model": model, "dataset": config.get("data", coco.get("data")), "seed": config.get("seed"),
        "training_status": state.get("status", "unknown"),
        "evaluation_status": "stale" if stale and evaluation.get("status", "completed") == "completed" else evaluation.get("status", "completed" if coco else "pending"),
        "stop_reason": state.get("stop_reason"),
        "epochs_completed": int(curves[-1]["epoch"]) if curves else state.get("epochs_completed"),
        "best_epoch": best_epoch, "best_epoch_source": best_source,
        "training_seconds": state.get("training_seconds"),
        "internal_precision": best_row.get("metrics/precision(B)"),
        "internal_recall": best_row.get("metrics/recall(B)"),
        "parameters": state.get("parameters"), "gflops": state.get("gflops"),
        "peak_gpu_memory_gib_rank0": state.get("peak_gpu_memory_gib_rank0"),
        "metrics": metrics, "per_class": coco.get("per_class", {}) if not stale else {},
        "protocol": {k: coco.get(k) for k in ("protocol", "maxDets", "annotations", "annotation_sha256", "inference")},
        "config": config, "files": file_paths,
        "files_exist": {k: Path(v).exists() for k, v in file_paths.items()},
        "curves": curves, "warnings": warnings,
        "evaluation_error": evaluation.get("error"),
    }


def comparable(a, b):
    # Unknown architecture/protocol/configuration is not evidence of equivalence.
    if not a["model"] or str(a["model"]).endswith(".pt") or a["model"] != b["model"]:
        return False
    if not a["dataset"] or a["dataset"] != b["dataset"]:
        return False
    if any(a["config"].get(k) is None or a["config"].get(k) != b["config"].get(k) for k in COMPARE_KEYS):
        return False
    pa, pb = a["protocol"], b["protocol"]
    return (pa.get("protocol") is not None and pa.get("maxDets") is not None
            and pa.get("annotations") is not None
            and all(pa.get(k) == pb.get(k) for k in pa))


def add_comparisons(records, baseline):
    for row in records:
        candidates = [b for b in records if b["mode"] == "fixed" and b["metrics"].get("AP") is not None
                      and (baseline is None or b["run"] == baseline)
                      and (b is row or comparable(row, b))]
        row["baseline"] = candidates[0]["run"] if len(candidates) == 1 else None
        row["delta_pp"] = {}
        if row["baseline"]:
            b = candidates[0]
            row["delta_pp"] = {k: (row["metrics"][k] - b["metrics"][k]) * 100 for k in COCO_STAT_NAMES
                               if row["metrics"].get(k) is not None and b["metrics"].get(k) is not None}
        elif row["metrics"]:
            row["warnings"].append("未找到唯一且配置、评估协议一致的 Q0；未计算提升")


def flat_rows(records):
    fields = ("run", "experiment", "description", "mode", "model", "dataset", "seed", "training_status",
              "evaluation_status", "stop_reason", "epochs_completed", "best_epoch", "best_epoch_source",
              "training_seconds", "internal_precision", "internal_recall", "parameters", "gflops",
              "peak_gpu_memory_gib_rank0", "baseline")
    return [{**{k: r[k] for k in fields}, **{k: r["metrics"].get(k) for k in COCO_STAT_NAMES},
             **{f"delta_{k}_pp": r["delta_pp"].get(k) for k in COCO_STAT_NAMES},
             "warnings": "；".join(r["warnings"])} for r in records]


def display(value, percent=False):
    return "—" if value is None else f"{value * (100 if percent else 1):.2f}"


def safe_text(value):
    # Prevent spreadsheet formula execution in user-supplied descriptions/paths.
    return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")) else value


def markdown(data, output):
    def escape(value):
        return str(value).replace("|", "\\|").replace("\n", "<br>")
    lines = ["# SO-DETR 实验汇总", "", f"更新时间：{data['generated_at']}", "",
             "COCO 指标以百分数显示，提升单位为百分点（pp）；JSON/CSV 保存 0–1 原始值。",
             "Precision/Recall 来自最佳轮次的内部验证日志，不属于正式 COCOeval 12 项指标。",
             "缺失或无有效 GT 的指标显示 —；历史训练状态不推测。不同配置的 AP 只供查看，不代表公平比较。", "",
             "## 主结果", "", "| 实验 / Run | 模式 | 评估状态 | AP | AP50 | AP75 | APs | APm | APl | ΔAP (pp) |",
             "|---|---|---|---:|---:|---:|---:|---:|---:|---:|"]
    for r in data["experiments"]:
        cells = [f"{r['experiment']} / {r['run']}", r["mode"], r["evaluation_status"]]
        cells += [display(r["metrics"].get(k), True) for k in COCO_STAT_NAMES[:6]]
        cells += [display(r["delta_pp"].get("AP"))]
        lines.append("| " + " | ".join(map(escape, cells)) + " |")
    lines += ["", "## 实验说明", ""]
    for r in data["experiments"]:
        lines += [f"### {escape(r['run'])}", "", escape(r["description"]), "",
                  f"训练状态：{r['training_status']}；停止原因：{r['stop_reason'] or '未知'}；"
                  f"训练轮数：{r['epochs_completed']}；最佳轮次：{r['best_epoch']}（{r['best_epoch_source']}）。", ""]
        if r["baseline"] and r["delta_pp"]:
            lines += [f"相对 {escape(r['baseline'])}：" + "，".join(
                f"{k} {r['delta_pp'][k]:+.2f} pp" for k in ("AP", "APs", "APm", "APl") if k in r["delta_pp"]) + "。", ""]
        if r["warnings"]:
            lines += ["数据说明：" + "；".join(map(escape, r["warnings"])) + "。", ""]
        if r["evaluation_error"]:
            lines += ["评估失败：" + escape(r["evaluation_error"]), ""]
        lines += [f"- [{key}]({Path(path).as_uri()})" + ("（文件不存在）" if not r["files_exist"][key] else "")
                  for key, path in r["files"].items()]
        lines += [""]
    lines += ["## 图表", ""]
    for name in ("map_comparison", "small_medium_large", "learning_curves"):
        lines += [f"![{name}](plots/{name}.png)", ""]
    if data["warnings"]:
        lines += ["## 扫描提示", ""] + [f"- {escape(w)}" for w in data["warnings"]]
    (output / "experiments.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def plots(records, output):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    target = output / "plots"
    target.mkdir(exist_ok=True)
    for name, keys in (("map_comparison", ("AP", "AP50", "AP75")),
                       ("small_medium_large", ("APs", "APm", "APl"))):
        fig, ax = plt.subplots(figsize=(max(8, len(records) * 1.4), 4.8))
        for j, k in enumerate(keys):
            ax.bar([i + (j - 1) * .25 for i in range(len(records))],
                   [r["metrics"].get(k) * 100 if r["metrics"].get(k) is not None else float("nan") for r in records],
                   width=.25, label=k)
        ax.set_xticks(range(len(records)), [r["run"] for r in records], rotation=20, ha="right")
        ax.set_ylabel("COCO AP (%)")
        ax.legend()
        ax.grid(axis="y", alpha=.2)
        fig.tight_layout()
        fig.savefig(target / f"{name}.png", dpi=150)
        plt.close(fig)
    fig, axes = plt.subplots(2, 3, figsize=(15, 8))
    keys = ("metrics/mAP50-95(B)", "metrics/precision(B)", "metrics/recall(B)",
            "train/giou_loss", "train/cls_loss", "train/l1_loss")
    for ax, key in zip(axes.flat, keys):
        for r in records:
            points = [p for p in r["curves"] if p.get(key) is not None]
            if points:
                ax.plot([p["epoch"] for p in points], [p[key] * (100 if key.startswith("metrics/") else 1) for p in points], label=r["run"])
        ax.set_title(key + " (internal)" if key.startswith("metrics/") else key)
        ax.set_xlabel("Epoch")
        ax.set_ylabel("%" if key.startswith("metrics/") else "Loss")
        ax.grid(alpha=.2)
        if ax.lines:
            ax.legend(fontsize=7)
    fig.tight_layout()
    fig.savefig(target / "learning_curves.png", dpi=150)
    plt.close(fig)


def excel(data, output):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import CellIsRule

    records = data["experiments"]
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([safe_text(v) for v in row])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="25476A")
        for i, col in enumerate(ws.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(55, max(14, max(len(str(c.value or "")) for c in col) + 2))
        return ws

    headers = ["实验", "Run", "说明", "模式", "训练状态", "评估状态", "AP (%)", "AP50 (%)", "AP75 (%)", "APs (%)", "APm (%)", "APl (%)", "ΔAP (pp)", "基线", "内部 P (%)", "内部 R (%)", "轮数", "最佳轮次", "耗时 (h)", "参数量", "GFLOPs", "峰值显存 GiB (rank0)", "数据说明"]
    overview = sheet("总览", headers, [[r["experiment"], r["run"], r["description"], r["mode"], r["training_status"], r["evaluation_status"],
        *[r["metrics"].get(k) * 100 if r["metrics"].get(k) is not None else None for k in COCO_STAT_NAMES[:6]],
        r["delta_pp"].get("AP"), r["baseline"],
        *[r[k] * 100 if r[k] is not None else None for k in ("internal_precision", "internal_recall")],
        r["epochs_completed"], r["best_epoch"], r["training_seconds"] / 3600 if r["training_seconds"] is not None else None,
        r["parameters"], r["gflops"], r["peak_gpu_memory_gib_rank0"], "；".join(r["warnings"])] for r in records])
    for row in overview.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
    if records:
        for op, color in (("greaterThan", "C6EFCE"), ("lessThan", "FFC7CE")):
            overview.conditional_formatting.add(f"M2:M{len(records)+1}", CellIsRule(operator=op, formula=["0"], fill=PatternFill("solid", fgColor=color)))
        # Highlight the best only inside each comparable baseline group.
        groups = {r["baseline"] for r in records if r["baseline"]}
        for group in groups:
            eligible = [(i + 2, r["metrics"].get("AP")) for i, r in enumerate(records) if r["baseline"] == group and r["metrics"].get("AP") is not None]
            top = max(v for _, v in eligible)
            for i, value in eligible:
                if value == top:
                    overview.cell(i, 7).fill = PatternFill("solid", fgColor="FFEB9C")
                    overview.cell(i, 7).font = Font(bold=True)
        chart = BarChart()
        chart.title = "COCO AP (%)"
        chart.add_data(Reference(overview, min_col=7, max_col=9, min_row=1, max_row=len(records)+1), titles_from_data=True)
        chart.set_categories(Reference(overview, min_col=2, min_row=2, max_row=len(records)+1))
        overview.add_chart(chart, f"A{len(records)+5}")
    sheet("COCO指标", ["Run", *[k + " (%)" for k in COCO_STAT_NAMES]],
          [[r["run"], *[r["metrics"].get(k) * 100 if r["metrics"].get(k) is not None else None for k in COCO_STAT_NAMES]] for r in records])
    sheet("PerClass AP", ["Run", "类别 ID", "类别", "AP (%)", "AP50 (%)", "AP75 (%)"],
          [[r["run"], cid, c["name"], *[c.get(k) * 100 if c.get(k) is not None else None for k in ("AP", "AP50", "AP75")]] for r in records for cid, c in r["per_class"].items()])
    curve_keys = sorted({k for r in records for p in r["curves"] for k in p if k != "epoch"})
    sheet("训练曲线", ["Run", "epoch", *curve_keys], [[r["run"], p["epoch"], *[p.get(k) for k in curve_keys]] for r in records for p in r["curves"]])
    config_keys = sorted({k for r in records for k in r["config"]})
    sheet("实验配置", ["Run", *config_keys], [[r["run"], *[json.dumps(r["config"].get(k), ensure_ascii=False) if isinstance(r["config"].get(k), (list, dict)) else r["config"].get(k) for k in config_keys]] for r in records])
    files = sheet("文件索引", ["Run", "文件", "路径", "存在"], [[r["run"], k, v, r["files_exist"][k]] for r in records for k, v in r["files"].items()])
    for row in files.iter_rows(min_row=2):
        row[2].hyperlink = Path(row[2].value).as_uri()
        row[2].style = "Hyperlink"
    gallery = wb.create_sheet("图表")
    for index, name in enumerate(("map_comparison", "small_medium_large", "learning_curves")):
        img = Image(str(output / "plots" / f"{name}.png"))
        img.width, img.height = 1000, 540
        gallery.add_image(img, f"A{index * 30 + 1}")
    sheet("口径说明", ["项目", "说明"], [["指标单位", "COCO/内部 P/R 在 Excel 为百分数；JSON/CSV 为 0–1；差值为百分点 pp"],
        ["内部 P/R", "训练日志最佳轮次，与正式 COCOeval 分开；CSV 估算轮次见总览提示"],
        ["比较", "仅唯一且配置/协议一致的 fixed 基线计算差值；黄色为同组最佳 AP，不跨组排名"],
        ["未知数据", "空白代表未记录或无有效 GT；GFLOPs 是模型工具估计；显存为 rank0 PyTorch 峰值 allocated"],
        ["更新时间", data["generated_at"]], *[["扫描提示", w] for w in data["warnings"]]])
    wb.save(output / "experiments.xlsx")


def generate_reports(runs_dir: Path, output: Path, baseline=None, completed_only: bool = False) -> dict:
    import fcntl
    output = output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    # Serialize writers and scan inside the lock to avoid losing concurrent runs.
    with (output / ".reports.lock").open("w") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        records, warnings = [], []
        candidates = sorted(p for p in runs_dir.resolve().iterdir() if p.is_dir()) if runs_dir.exists() else []
        for run in candidates:
            if not any((run / p).exists() for p in ("args.yaml", "results.csv", "formal_coco/best/coco_metrics.json")):
                continue
            try:
                rec = collect_run(run)
                if completed_only and (rec.get("evaluation_status") != "completed" or not rec.get("metrics")):
                    continue
                records.append(rec)
            except (ValueError, OSError, TypeError, AttributeError, yaml.YAMLError) as exc:
                warnings.append(f"跳过损坏实验 {run.name}: {exc}")
        add_comparisons(records, baseline)
        data = {"schema_version": 1, "generated_at": datetime.now(timezone.utc).isoformat(),
                "units": {"metrics": "ratio (0–1)", "delta_pp": "percentage points"},
                "experiments": records, "warnings": warnings}
        # Every renderer consumes the same serialized snapshot.
        with tempfile.TemporaryDirectory(dir=output, prefix=".report-") as temporary:
            stage = Path(temporary)
            write_json(stage / "experiments.json", data)
            data = read_json(stage / "experiments.json")
            flattened = flat_rows(data["experiments"])
            with (stage / "experiments.csv").open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=list(flattened[0]) if flattened else ["run", "AP"])
                writer.writeheader()
                writer.writerows({k: safe_text(v) for k, v in row.items()} for row in flattened)
            plots(data["experiments"], stage)
            markdown(data, stage)
            excel(data, stage)
            (output / "plots").mkdir(exist_ok=True)
            for path in stage.rglob("*"):
                if path.is_file():
                    os.replace(path, output / path.relative_to(stage))
        return data


def add_report_arguments(parser):
    parser.add_argument("--reports", action=argparse.BooleanOptionalAction, default=True,
                        help="Automatically refresh JSON/CSV/Markdown/Excel reports after training.")
    parser.add_argument("--reports-dir", type=Path, default=ROOT / "reports")
    parser.add_argument("--baseline", default=None, help="Exact baseline run directory name; default: unique compatible fixed run.")
    parser.add_argument("--experiment-description", default=None, help="Human-readable experiment description.")


def backfill_checkpoint_info(checkpoint: Path):
    """Explicit opt-in only: unpickle trusted local training checkpoints."""
    import torch
    run = checkpoint.parent.parent
    state = read_json(run / "training_state.json")
    checkpoint_data = torch.load(checkpoint, map_location="cpu", weights_only=False)
    model = checkpoint_data.get("ema") or checkpoint_data.get("model")
    state.setdefault("status", "unknown")
    if model is not None:
        state["model"] = model.yaml.get("yaml_file")
        state["parameters"] = sum(p.numel() for p in model.parameters())
    epoch = checkpoint_data.get("epoch", -1)
    if epoch < 0:
        epochs = checkpoint_data.get("train_results", {}).get("epoch", [])
        epoch = int(epochs[-1]) - 1 if epochs else -1
    if epoch >= 0 and state.get("best_epoch") is None:
        state.update(best_epoch=epoch + 1, best_epoch_source="best.pt checkpoint")
    state["metadata_source"] = str(checkpoint.resolve())
    write_json(run / "training_state.json", state)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--runs-dir", type=Path, default=ROOT / "runs/train")
    parser.add_argument("--output", type=Path, default=ROOT / "reports")
    parser.add_argument("--baseline")
    parser.add_argument("--evaluate", action="append", default=[], metavar="RUN", help="Re-run best.pt inference and formal COCOeval for a named run, then refresh reports.")
    parser.add_argument("--coco-anno", type=Path)
    parser.add_argument("--device", help="Override device when using --evaluate.")
    parser.add_argument("--backfill-checkpoint-info", action="store_true", help="Read trusted local best.pt files to recover model, parameters and best epoch. Requires the training environment.")
    parser.add_argument("--backfill-per-class", action="store_true", help="CPU COCOeval of saved normalized predictions for old runs missing per-class AP.")
    parser.add_argument("--completed-only", action="store_true", help="Only include runs with completed formal evaluation, excluding pending or interrupted runs.")
    args = parser.parse_args()
    failures = []
    for name in args.evaluate:
        run = (args.runs_dir / name).resolve()
        if run.parent != args.runs_dir.resolve():
            parser.error("--evaluate must be a direct run directory name")
        try:
            from sodetr_formal_coco import prepare_formal_coco_eval, run_formal_coco_eval
            cfg = yaml.safe_load((run / "args.yaml").read_text())
            write_json(run / "evaluation_state.json", {"status": "running"})
            old = read_json(run / "formal_coco/best/coco_metrics.json")
            anno = prepare_formal_coco_eval(cfg["data"], args.coco_anno or old.get("annotations"), True)
            run_formal_coco_eval(cfg["data"], anno, run, cfg["imgsz"], cfg["batch"], cfg["workers"], args.device or cfg["device"])
            write_json(run / "evaluation_state.json", {"status": "completed"})
        except Exception as exc:
            write_json(run / "evaluation_state.json", {"status": "failed", "error": str(exc)})
            failures.append(f"{name}: {exc}")
            print(f"[Evaluation failed] {failures[-1]}", flush=True)
    if args.backfill_checkpoint_info:
        for checkpoint in sorted(args.runs_dir.glob("*/weights/best.pt")):
            try:
                backfill_checkpoint_info(checkpoint)
            except Exception as exc:
                failures.append(f"{checkpoint}: {exc}")
                print(f"[Checkpoint metadata failed] {failures[-1]}", flush=True)
    if args.backfill_per_class:
        for path in sorted(args.runs_dir.glob("*/formal_coco/best/coco_metrics.json")):
            try:
                if not read_json(path).get("per_class"):
                    enrich_saved_evaluation(path)
            except Exception as exc:
                failures.append(f"{path}: {exc}")
                print(f"[Backfill failed] {failures[-1]}", flush=True)
    data = generate_reports(args.runs_dir, args.output, args.baseline, completed_only=args.completed_only)
    print(f"[Reports] {len(data['experiments'])} experiments → {args.output.resolve()}")
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
