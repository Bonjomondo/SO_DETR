"""Regression tests for COCO extraction, report provenance and exports (CPU only)."""
import copy
import csv
import json
import tempfile
import unittest
from pathlib import Path

import yaml
from openpyxl import load_workbook

from sodetr_formal_coco import _normalize_predictions, evaluate_predictions
from sodetr_reports import (COMPARE_KEYS, add_comparisons, collect_run, generate_reports,
                            read_json, write_json)


class ExperimentReportsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)

    def fixture(self, name='Q0', mode='fixed', ap=.3):
        run = self.root / 'runs' / name
        run.mkdir(parents=True)
        config = {k: 1 for k in COMPARE_KEYS}
        config.update(model='r18.yaml', data='val.yaml', expanded_iou_mode=mode)
        (run / 'args.yaml').write_text(yaml.safe_dump(config))
        (run / 'results.csv').write_text(' epoch, metrics/precision(B), metrics/recall(B), metrics/mAP50(B), metrics/mAP50-95(B)\n1,.8,.4,.6,.3\n2,.9,.5,.5,.29\n')
        write_json(run / 'formal_coco/best/coco_metrics.json', {
            'metrics': {'AP': ap, 'AP50': .5, 'APs': -1},
            'protocol': 'pycocotools.COCOeval bbox', 'maxDets': [1, 10, 100], 'annotations': 'gt.json',
            'per_class': {'1': {'name': 'car', 'AP': ap}}})
        return run

    def test_legacy_best_epoch_missing_and_stale(self):
        run = self.fixture()
        row = collect_run(run)
        self.assertEqual(row['best_epoch'], 1)
        self.assertEqual(row['internal_precision'], .8)
        self.assertEqual(row['training_status'], 'unknown')
        self.assertIsNone(row['training_seconds'])
        self.assertIsNone(row['metrics']['APs'])
        write_json(run / 'evaluation_state.json', {'status': 'failed', 'error': 'test error'})
        self.assertEqual(collect_run(run)['metrics'], {})

    def test_comparison_units_and_incompatible_baselines(self):
        q0 = collect_run(self.fixture())
        q1 = collect_run(self.fixture('Q1', 'adaptive-quality', .318))
        add_comparisons([q0, q1], None)
        self.assertAlmostEqual(q1['delta_pp']['AP'], 1.8)
        mismatch = copy.deepcopy(q1)
        mismatch['config']['seed'] = 2
        add_comparisons([q0, mismatch], None)
        self.assertIsNone(mismatch['baseline'])
        duplicate = copy.deepcopy(q0)
        duplicate['run'] = 'Q0-repeat'
        add_comparisons([q0, duplicate, q1], None)
        self.assertIsNone(q1['baseline'])
        add_comparisons([q0, duplicate, q1], 'Q0')
        self.assertEqual(q1['baseline'], 'Q0')
        q1['protocol']['maxDets'] = [1, 10, 300]
        add_comparisons([q0, q1], 'Q0')
        self.assertIsNone(q1['baseline'])

    def test_report_exports_and_corrupt_run(self):
        run = self.fixture()
        self.fixture('Q1', 'adaptive-quality', .318)
        write_json(run / 'experiment.json', {'description': '=SUM(1,2)'})
        bad = self.root / 'runs' / 'broken'
        bad.mkdir()
        (bad / 'args.yaml').write_text('not: [valid')
        data = generate_reports(self.root / 'runs', self.root / 'reports')
        self.assertEqual(len(data['experiments']), 2)
        self.assertEqual(len(data['warnings']), 1)
        output = self.root / 'reports'
        self.assertEqual(len(read_json(output / 'experiments.json')['experiments']), 2)
        with (output / 'experiments.csv').open(encoding='utf-8-sig') as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual(rows[0]['AP'], '0.3')
        self.assertEqual(rows[0]['description'], "'=SUM(1,2)")
        wb = load_workbook(output / 'experiments.xlsx')
        self.assertEqual(wb['总览']['G2'].value, 30)
        self.assertAlmostEqual(wb['总览']['M3'].value, 1.8)
        self.assertEqual(wb['总览']['C2'].data_type, 's')
        self.assertEqual(wb['总览'].freeze_panes, 'A2')
        self.assertIn('PerClass AP', wb.sheetnames)
        self.assertEqual(len(list((output / 'plots').glob('*.png'))), 3)
        self.assertIn('30.00', (output / 'experiments.md').read_text())
        # Repeated generation replaces outputs rather than appending duplicate rows.
        again = generate_reports(self.root / 'runs', output)
        self.assertEqual(len(again['experiments']), 2)

    def test_completed_only_filtering(self):
        self.fixture('Q0', 'fixed', .30)
        interrupted = self.root / 'runs' / 'Q1_interrupted'
        interrupted.mkdir(parents=True, exist_ok=True)
        (interrupted / 'args.yaml').write_text('expanded_iou_mode: adaptive-quality\n')
        (interrupted / 'results.csv').write_text('epoch,train/giou_loss\n1,0.5\n')
        write_json(interrupted / 'evaluation_state.json', {'status': 'pending'})

        all_data = generate_reports(self.root / 'runs', self.root / 'reports', completed_only=False)
        self.assertEqual(len(all_data['experiments']), 2)

        comp_data = generate_reports(self.root / 'runs', self.root / 'reports', completed_only=True)
        self.assertEqual(len(comp_data['experiments']), 1)
        self.assertEqual(comp_data['experiments'][0]['run'], 'Q0')

    def test_empty_project_exports(self):
        data = generate_reports(self.root / 'missing', self.root / 'reports')
        self.assertEqual(data['experiments'], [])
        self.assertTrue((self.root / 'reports/experiments.xlsx').is_file())

    def test_coco_per_class_and_empty_predictions(self):
        gt = {'info': {}, 'images': [{'id': 1, 'file_name': 'one.jpg', 'width': 100, 'height': 100}],
              'categories': [{'id': 1, 'name': 'car'}, {'id': 2, 'name': 'bus'}],
              'annotations': [{'id': 1, 'image_id': 1, 'category_id': 1, 'bbox': [10, 10, 20, 20],
                               'area': 400, 'iscrowd': 0}]}
        annotation = self.root / 'gt.json'
        predictions = self.root / 'pred.json'
        write_json(annotation, gt)
        write_json(predictions, [{'image_id': 1, 'category_id': 1, 'bbox': [10, 10, 20, 20], 'score': .99}])
        result = evaluate_predictions(annotation, predictions)
        self.assertAlmostEqual(result['metrics']['AP'], 1)
        self.assertAlmostEqual(result['per_class']['1']['AP'], 1)
        self.assertIsNone(result['per_class']['2']['AP'])
        self.assertIsNone(result['metrics']['APl'])
        write_json(predictions, [])
        result = evaluate_predictions(annotation, predictions)
        self.assertEqual(result['metrics']['AP'], 0)
        self.assertEqual(result['per_class']['1']['AP'], 0)
        self.assertIsNone(result['per_class']['2']['AP'])

    def test_category_indices_overlapping_gt_still_map_by_name(self):
        annotation, raw, data, output = [self.root / p for p in ('gt.json', 'raw.json', 'data.yaml', 'out.json')]
        write_json(annotation, {'images': [{'id': 42, 'file_name': 'one.jpg'}],
                               'categories': [{'id': 1, 'name': 'car'}, {'id': 2, 'name': 'bus'}]})
        data.write_text('names: [car, bus]\n')
        write_json(raw, [{'image_id': 'one', 'category_id': 1, 'bbox': [0, 0, 1, 1], 'score': .5}])
        _normalize_predictions(raw, annotation, data, output)
        prediction = json.loads(output.read_text())[0]
        self.assertEqual(prediction['image_id'], 42)
        self.assertEqual(prediction['category_id'], 2)


if __name__ == '__main__':
    unittest.main()
